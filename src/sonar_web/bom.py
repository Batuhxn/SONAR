"""Bounded, loss-aware engineering BoM import and validation."""
from __future__ import annotations

import csv
import io
import re
import uuid
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

MAX_ROWS = 5000
MAX_COLUMNS = 100
MAX_CELL = 2000
FIELDS = ("name", "mpn", "quantity", "references", "dnp")
ALIASES = {
    "name": {"name", "component", "componentname", "description", "comment", "value", "part"},
    "mpn": {"mpn", "manufacturerpartnumber", "manufacturerpart", "manufacturerspartnumber", "partnumber", "mfrpartnumber", "mfrpn"},
    "quantity": {"quantity", "qty", "count", "partcount"},
    "references": {"references", "reference", "referencedesignators", "referencedesignator", "designator", "designators", "ref", "refs", "refdes"},
    "dnp": {"dnp", "dnf", "donotpopulate", "donotfit", "notfitted", "fitted", "populate", "populated", "exclude from bom"},
}


def header_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


def recognize(headers: list[str]) -> dict[str, int]:
    result = {}
    for field_name, aliases in ALIASES.items():
        matches = [i for i, h in enumerate(headers) if header_key(h) in {header_key(a) for a in aliases}]
        if matches:
            result[field_name] = matches[0]
    return result


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_table(data: bytes, filename: str, sheet: str | None = None) -> dict:
    suffix = Path(filename).suffix.lower()
    sheets = []
    if suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 1000 or sum(e.file_size for e in entries) > 30_000_000:
                    raise ValueError("Expanded workbook exceeds the 30 MB / 1000-entry limit")
                if any(e.flag_bits & 1 for e in entries):
                    raise ValueError("Encrypted workbooks are unsupported")
                if any("vbaproject" in e.filename.lower() for e in entries):
                    raise ValueError("Macro workbooks are unsupported")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                sheets = workbook.sheetnames
                if sheet and sheet not in sheets:
                    raise ValueError("Worksheet not found")
                selected = workbook[sheet or sheets[0]]
                # Ignore untrusted worksheet dimension hints; enforce bounds while reading.
                selected.reset_dimensions()
                rows = []
                for row in selected.iter_rows(values_only=True):
                    if len(row) > MAX_COLUMNS or any(len(text(cell)) > MAX_CELL for cell in row):
                        raise ValueError("Worksheet exceeds the column / cell size limit")
                    rows.append(list(row))
                    if len(rows) > MAX_ROWS + 50:
                        raise ValueError("BoM exceeds the 5000 component row limit")
            finally:
                workbook.close()
        except ValueError:
            raise
        except Exception as error:
            raise ValueError("Invalid or unsupported XLSX workbook") from error
    elif suffix in {".csv", ".tsv"}:
        if b"\x00" in data:
            raise ValueError("Binary or UTF-16 input is unsupported; save as UTF-8 CSV/TSV")
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            # Common Turkish Windows export encoding; original values remain visible.
            content = data.decode("cp1254")
        try:
            delimiter = "\t" if suffix == ".tsv" else csv.Sniffer().sniff(content[:16000], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ","
        try:
            rows = []
            for row in csv.reader(io.StringIO(content), delimiter=delimiter, strict=True):
                rows.append(row)
                if len(rows) > MAX_ROWS + 50:
                    raise ValueError("BoM exceeds the 5000 component row limit")
        except csv.Error as error:
            raise ValueError("Malformed CSV/TSV quoting or cell size") from error
    else:
        raise ValueError("Use Altium CSV/XLSX or KiCad CSV/TSV")
    if not rows or not any(any(text(c) for c in r) for r in rows):
        raise ValueError("The BoM is empty")
    if any(len(row) > MAX_COLUMNS for row in rows):
        raise ValueError("BoM exceeds the 100-column limit")
    if any(len(text(cell)) > MAX_CELL for row in rows for cell in row):
        raise ValueError("A cell exceeds the 2000-character limit")
    # Find headers after common engineering report preambles.
    header_row = max(range(min(len(rows), 50)), key=lambda i: len(recognize([text(c) for c in rows[i]])))
    width = max(len(row) for row in rows[header_row:])
    headers = [text(c) or f"Column {i + 1}" for i, c in enumerate(rows[header_row])]
    headers += [f"Extra column {i + 1}" for i in range(len(headers), width)]
    body = [["" if c is None else str(c) for c in row] + [""] * (width - len(row)) for row in rows[header_row + 1:]]
    # Only completely blank physical rows are omitted; malformed nonblank rows survive.
    body = [row for row in body if any(text(c) for c in row)]
    if not body:
        raise ValueError("No component rows found after the header")
    if len(body) > MAX_ROWS:
        raise ValueError("BoM exceeds the 5000 component row limit")
    return {"headers": headers, "rows": body, "mapping": recognize(headers), "header_row": header_row + 1,
            "preamble": [[text(c) for c in row] for row in rows[:header_row]],
            "sheets": sheets, "sheet": sheet or (sheets[0] if sheets else None)}


def parse_quantity(raw: str) -> int | None:
    try:
        number = Decimal(raw)
        if number.is_finite() and number == number.to_integral_value() and 0 < number <= 1_000_000:
            return int(number)
    except (InvalidOperation, ValueError):
        pass
    return None


def parse_dnp(raw: str, inverted: bool = False) -> bool | None:
    key = header_key(raw)
    if inverted:
        if key in {"1", "true", "yes", "y", "fitted", "populated", "populate"}:
            return False
        if key in {"0", "false", "no", "n", "dnp", "dnf", "notfitted", "excluded"}:
            return True
        return None
    if key in {"1", "true", "yes", "y", "dnp", "dnf", "donotpopulate", "notfitted", "excluded"}:
        result = True
    elif key in {"", "0", "false", "no", "n", "fitted", "populated"}:
        result = False
    else:
        return None
    return result


def references(raw: str) -> list[str]:
    result = []
    for token in re.split(r"[,;\s]+", raw.strip()):
        if not token:
            continue
        match = re.fullmatch(r"([A-Za-z]+)(\d+)-([A-Za-z]*)(\d+)", token)
        if match and (not match[3] or match[1].upper() == match[3].upper()):
            start, end = int(match[2]), int(match[4])
            if 0 <= end - start and len(result) + end - start + 1 <= MAX_ROWS:
                result.extend(f"{match[1].upper()}{i}" for i in range(start, end + 1))
                continue
        result.append(token.upper())
    return result


@dataclass
class Component:
    id: str
    name: str = ""
    mpn: str = ""
    quantity: str = ""
    references: str = ""
    dnp: bool | None = False
    raw: dict = field(default_factory=dict)
    dnp_raw: str = ""
    selected: bool = True
    issues: list[str] = field(default_factory=list)
    required: int | None = None
    results: dict = field(default_factory=dict)
    choice: dict | None = None
    revision: int = 0

    def to_dict(self):
        return asdict(self)


def import_components(table: dict, mapping: dict[str, int]) -> list[Component]:
    if any(k not in FIELDS or type(v) is not int or not 0 <= v < len(table["headers"]) for k, v in mapping.items()):
        raise ValueError("Invalid column mapping")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("Each source column can map to only one field")
    if not {"name", "mpn", "references"}.intersection(mapping):
        raise ValueError("Map at least a component name, MPN or reference column")
    components = []
    for row in table["rows"]:
        values = {key: text(row[index]) for key, index in mapping.items()}
        ref = values.get("references", "")
        qty = values.get("quantity", "")
        if "quantity" not in mapping and ref:
            qty = str(len(references(ref)))
        dnp_raw = values.get("dnp", "")
        inverted = "dnp" in mapping and header_key(table["headers"][mapping["dnp"]]) in {"fitted", "populate", "populated"}
        dnp = parse_dnp(dnp_raw, inverted)
        component = Component(uuid.uuid4().hex, values.get("name", ""), values.get("mpn", "").strip().upper(), qty, ref,
                              dnp, {f"{i + 1}: {h}": row[i] for i, h in enumerate(table["headers"])}, dnp_raw,
                              selected=dnp is not True)
        components.append(component)
    validate(components, 1)
    return components


def validate(components: list[Component], boards: int) -> None:
    if type(boards) is not int or not 1 <= boards <= 100000:
        raise ValueError("PCB quantity must be an integer from 1 to 100000")
    ref_counts = Counter(ref for c in components for ref in references(c.references))
    keys = Counter((c.mpn, c.name.casefold(), c.references.upper()) for c in components)
    for c in components:
        c.issues = []
        quantity = parse_quantity(c.quantity)
        c.required = 0 if c.dnp is True else (quantity * boards if quantity is not None and c.dnp is False else None)
        if quantity is None:
            c.issues.append("Quantity must be a positive whole number (maximum 1000000)")
        if c.dnp is None:
            c.issues.append("Unrecognized DNP value; choose Populate or DNP")
        if not c.name and not c.mpn:
            c.issues.append("Component name and MPN are both missing")
        if not c.mpn:
            c.issues.append("Missing MPN; use a manual URL and verify the product")
        if not c.references:
            c.issues.append("Reference designators missing")
        refs = references(c.references)
        if refs and quantity is not None and len(refs) != quantity:
            c.issues.append("Reference count differs from quantity; review the source")
        if any(not re.fullmatch(r"[A-Z]+\d+", r) for r in refs):
            c.issues.append("Unrecognized reference format")
        if any(ref_counts[r] > 1 for r in refs):
            c.issues.append("Duplicate reference designator; rows were kept separately")
        if keys[(c.mpn, c.name.casefold(), c.references.upper())] > 1:
            c.issues.append("Possible duplicate row; no automatic merging")
        if any(v.lstrip().startswith("=") for v in c.raw.values()):
            c.issues.append("Source contains a formula; formulas are preserved as text, never evaluated")
